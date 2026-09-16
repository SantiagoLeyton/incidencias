"""Lector sin escritura de libros históricos y detector de incidencias."""
from __future__ import annotations
import re
import calendar
from collections import Counter, defaultdict
from dataclasses import asdict
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook
from date_parser import parse_date

LABELS = {
 "HIST_INVALID_DATE_HEADER": "Fecha o encabezado de fecha inválido",
 "HIST_PAYMENT_DATE_RECEIPT_MISMATCH": "No coincide la cantidad de fechas y recibos",
 "HIST_PAYMENT_VALUE_COUNT_MISMATCH": "No coincide la cantidad de valores de pago",
 "STRUCTURE": "Estructura de hoja incompleta",
}
MONTHS = {"ENE":1,"FEB":2,"MAR":3,"ABR":4,"MAY":5,"JUN":6,"JUL":7,"AGO":8,"SEP":9,"OCT":10,"NOV":11,"DIC":12}

def normal(s): return re.sub(r"\s+", " ", str(s or "").strip().upper())
def cell_text(v):
    if v is None: return ""
    if isinstance(v, (datetime, date)): return v.strftime("%d/%m/%Y")
    return str(v).strip()

def header_kind(value):
    v = normal(value)
    if re.fullmatch(r"RECIBO FIDUCIA (?:[A-Z]{3}|SEPTIEMBRE|OCTUBRE|NOVIEMBRE|DICIEMBRE|ENERO|FEBRERO|MARZO|ABRIL|MAYO|JUNIO|JULIO|AGOSTO)[ /-]\d{2,4}", v): return "fid_value"
    if v == "RECIBOS FIDUBOGOTA": return "fid_receipts"
    if v == "RECIBOS": return "receipts"
    if v == "FECHA": return "dates"
    if v == "RECIBIDO": return "received"
    return None

def fiduciary_period(header):
    found = re.search(r"([A-Z]{3})[ /-](\d{2,4})$", normal(header))
    if not found: return None
    m, y = found.groups()
    return (MONTHS.get(m), 2000+int(y) if len(y)==2 else int(y))

def split_receipts(text):
    return [x.strip() for x in re.split(r"[\n;,-]+", text) if x.strip()]
def split_dates(text):
    # El guión se acepta únicamente antes de otro mes, no como partición genérica.
    return [x.strip(" ()") for x in re.split(r"[\n;,]+|-(?=[A-Za-z]{3,})", text) if x.strip(" ()")]
def split_values(text):
    text = text.strip()
    if not text or text in ("0", "0.0", "0,0", "."): return []
    # Fórmulas de suma se descomponen para preservar posiciones; no se evalúan.
    if text.startswith("="):
        return [x for x in re.findall(r"(?<![A-Z0-9_])-?\d+(?:[.,]\d+)?", text[1:])]
    return [x.strip() for x in re.split(r"[\n;,]+", text) if x.strip()]

# Adaptación autocontenida de las reglas blocking de
# fiduciary.imports.historical.parser. No se importa Gestión Fiduciaria.
STRICT_DATE_RE = re.compile(r"^(ENE|FEB|MAR|ABR|MAY|JUN|JUL|AGO|SEP|OCT|NOV|DIC)\.(\d{1,2})/(\d{2})(F)?$", re.I)
def gf_dates(value):
    text = cell_text(value)
    if not text: return []
    text = text.partition("||")[0].strip()
    if re.match(r"^\d{4}-\d{2}-\d{2}(?:\s+\d{2}:\d{2}:\d{2})?$", text): return [text]
    return [x.strip() for x in re.split(r"\s*-\s*", text) if x.strip()]
def gf_date_base(value):
    return re.sub(r"(TRASLADO|TRASL|CESION)$", "", str(value).strip().strip("()").strip(), flags=re.I).strip()
def gf_looks_date(value):
    text=gf_date_base(value); compact=re.sub(r"\s+","",text).upper()
    return bool(text and not re.match(r"^\d{4}-\d{2}-\d{2}",compact) and re.match(r"^[A-Z]{3,4}",compact) and any(c.isdigit() for c in compact) and any(c in compact for c in (".","/","-","F")))
def gf_valid_date(value):
    match=STRICT_DATE_RE.match(gf_date_base(value))
    if not match: return False
    month=MONTHS[match.group(1).upper()]; day=int(match.group(2)); year=2000+int(match.group(3))
    return 1 <= day <= calendar.monthrange(year,month)[1]
def gf_category_receipt(value):
    text=str(value).strip(); upper=text.upper(); compact=re.sub(r"[\s._/()\-]+","",text).upper()
    if re.search(r"\d\s*[\s._/-]*TRASL(?:ADO)?(?:[\s._/-]|$)",upper) or re.search(r"\dTRASLADO$",compact): return "transfer"
    if re.search(r"\d\s*[\s._/-]*CESION(?:[\s._/-]|$)",upper) or re.search(r"\dCESION$",compact): return "cession"
    if re.search(r"\d(?:SUB|SB)$",compact): return "subsidy"
    if re.search(r"\dCR$",compact): return "credit"
    return "ordinary"
def gf_category_date(value):
    compact=re.sub(r"[\s._/()\-]+","",str(value)).upper()
    if re.search(r"\d{2,4}TRASL(?:ADO)?$",compact): return "transfer"
    if re.search(r"\d{2,4}CESION$",compact): return "cession"
    return "ordinary"
def gf_receipts(value, source):
    tokens=[]
    for part in re.split(r"\s*-\s*",cell_text(value)):
        if not part.strip(): continue
        # La excepción de marcadores embebidos es la misma que usa GF.
        if re.search(r"\(?[A-Z]{3,}\.\d{1,2}/\d{2,4}(?:F|CESION|TRASL(?:ADO)?)?\)?",part.strip(),re.I):
            parts=[part.strip()]
        else: parts=[x.strip() for x in re.split(r"\s*/\s*",part) if x.strip()]
        tokens.extend({"value":x,"source":source,"category":gf_category_receipt(x)} for x in parts)
    return tokens
def gf_positive_values(value):
    text=cell_text(value)
    if not text: return []
    if text.startswith("="):
        expression=text[1:].replace(" ","")
        if not expression or re.search(r"[A-Za-z()*/^]",expression): return []
        if "-" in expression.lstrip("-"):
            parts=re.findall(r"[+-]?\d+(?:[.,]\d+)?",expression)
            if not parts or "".join(parts)!=expression: return []
            try: return [sum(Decimal(x.replace(",", ".")) for x in parts)]
            except InvalidOperation: return []
        parts=expression.split("+")
        if len(parts)<2 or any(not x for x in parts): return []
    else: parts=[text]
    values=[]
    for part in parts:
        try: amount=Decimal(part.replace("$","").replace(" ","").replace(",","."))
        except InvalidOperation: continue
        if amount>0: values.append(amount)
    return values
def gf_strict_parts(value):
    match=STRICT_DATE_RE.match(gf_date_base(value))
    if not match or not gf_valid_date(value): return None
    return 2000+int(match.group(3)), MONTHS[match.group(1).upper()], int(match.group(2))
def gf_replace_year(value, year):
    match=STRICT_DATE_RE.match(gf_date_base(value))
    if not match: return None
    suffix="F" if match.group(4) else ""
    candidate=f"{match.group(1).upper()}.{int(match.group(2))}/{year % 100:02d}{suffix}"
    return candidate if gf_valid_date(candidate) else None
def gf_normalize_fiduciary_years(pairs, values):
    # Misma corrección conservadora de GF: solo +/- un año, mismo mes de
    # columna y únicamente si mejora la secuencia cronológica vecina.
    normalized=list(pairs)
    for i, ((receipt,date_value),(amount,header)) in enumerate(zip(pairs,values)):
        parts=gf_strict_parts(date_value); period=fiduciary_period(header)
        if not parts or not period: continue
        year,month,_=parts
        if month!=period[0] or year==period[1] or abs(period[1]-year)!=1: continue
        candidate=gf_replace_year(date_value,period[1])
        if not candidate: continue
        candidate_parts=gf_strict_parts(candidate)
        previous=next((gf_strict_parts(x[1]) for x in reversed(normalized[:i]) if gf_strict_parts(x[1])),None)
        following=next((gf_strict_parts(x[1]) for x in normalized[i+1:] if gf_strict_parts(x[1])),None)
        def violations(parts):
            key=(parts[0],parts[1],parts[2]); return int(previous is not None and key <= previous)+int(following is not None and key >= following)
        if violations(candidate_parts)<violations(parts):
            normalized[i]=(receipt,candidate)
    return normalized
def gf_blocking(sheet, row, context, active, values, evidence):
    def value(key):
        return cell_text(values[active[key][0][0]-1]) if key in active and active[key][0][0]<=len(values) else ""
    date_values=gf_dates(value("dates"))
    invalid=[d for d in date_values if gf_looks_date(d) and not gf_valid_date(d)]
    if invalid:
        return [make_incident("HIST_INVALID_DATE_HEADER",sheet,row,context,
            "Fecha historica invalida segun el formato estricto de Gestión Fiduciaria.",
            [f"FECHA ORIGINAL: {d}"] + (
                [f"CORRECCION PROPUESTA: {parse_date(d).normalized}"] if parse_date(d).status=="RECUPERABLE" else []
            ),evidence) for d in invalid]
    receipts=gf_receipts(value("receipts"),"constructora")+gf_receipts(value("fid_receipts"),"fiduciaria")
    ordinary=[x for x in receipts if x["category"]=="ordinary"]
    ordinary_dates=[x for x in date_values if gf_category_date(x)=="ordinary"]
    has_fid="fid_receipts" in active
    standard=[x for x in ordinary if x["source"]=="constructora"]
    fiduciary=[x for x in ordinary if x["source"]=="fiduciaria"]
    c_dates=[x for x in ordinary_dates if not x.rstrip().rstrip(")").rstrip().upper().endswith("F")]
    f_dates=[x for x in ordinary_dates if x.rstrip().rstrip(")").rstrip().upper().endswith("F")]
    groups=([("fiduciaria",fiduciary,f_dates)] if has_fid and fiduciary and not standard else
            ([("constructora",standard,c_dates),("fiduciaria",fiduciary,f_dates)] if has_fid else [("legacy",ordinary,ordinary_dates)]))
    result=[]; pairs={}
    for destination, recs, dates in groups:
        if (recs or dates) and len(recs)!=len(dates):
            result.append(make_incident("HIST_PAYMENT_DATE_RECEIPT_MISMATCH",sheet,row,context,
                "La cantidad de fechas ordinarias no coincide con la cantidad de recibos ordinarios.",
                [f"DESTINO: {destination.upper()}",f"FECHAS: {len(dates)}",f"RECIBOS: {len(recs)}"],evidence))
        if len(recs)==len(dates): pairs[destination]=list(zip(recs,dates))
    received=gf_positive_values(value("received"))
    c_pairs=pairs.get("constructora",[])
    if c_pairs or received:
        if len(received)!=len(c_pairs):
            result.append(make_incident("HIST_PAYMENT_VALUE_COUNT_MISMATCH",sheet,row,context,
                "La cantidad de valores no permite reconstruir pagos históricos individuales.",
                [f"DESTINO: CONSTRUCTORA",f"VALORES: {len(received)}",f"RECIBOS CON FECHA: {len(c_pairs)}"],evidence))
    # GF evalúa Fiducia solo cuando sus recibos y fechas coinciden.
    f_pairs=pairs.get("fiduciaria",[])
    if len(fiduciary)==len(f_dates):
        f_values=[]
        for col, header in active.get("fid_value",[]):
            for amount in gf_positive_values(cell_text(values[col-1]) if col<=len(values) else ""):
                f_values.append((amount,header))
        if len(f_values)!=len(f_pairs):
            result.append(make_incident("HIST_PAYMENT_VALUE_COUNT_MISMATCH",sheet,row,context,
                "La cantidad de valores no permite reconstruir pagos históricos individuales.",
                [f"DESTINO: FIDUCIARIA",f"VALORES: {len(f_values)}",f"RECIBOS CON FECHA: {len(f_pairs)}"],evidence))
        elif f_pairs:
            f_pairs=gf_normalize_fiduciary_years(f_pairs,f_values)
            for (receipt, date_value),(amount,header) in zip(f_pairs,f_values):
                period=fiduciary_period(header); parsed=parse_date(date_value)
                if period and parsed.status=="VALIDA" and (parsed.year,parsed.month)!=(period[1],period[0]):
                    result.append(make_incident("HIST_PAYMENT_VALUE_COUNT_MISMATCH",sheet,row,context,
                        "El recibo, fecha y valor coinciden en cantidad, pero no en periodo.",
                        [f"RECIBO: {receipt['value']}",f"FECHA: {date_value}",f"COLUMNA DE ORIGEN: {header}"],evidence))
    return result

def context_marker(values, current):
    joined = " ".join(normal(v) for v in values if v is not None)
    if "CESION" in joined: return "CESI\u00d3N"
    if "TRASLADO" in joined: return "TRASLADO"
    return current

def make_incident(kind, sheet, row, context, description, info, evidence):
    return {"id": None, "kind":kind, "label":LABELS[kind], "sheet":sheet, "row":row,
            "context":context, "description":description, "info":info, "evidence":evidence}

def rows_for(source):
    max_len=max(len(source["dates"]),len(source["receipts"]),len(source["values"]),1)
    result=[]
    for n in range(max_len):
        d=source["dates"][n] if n<len(source["dates"]) else None
        r=source["receipts"][n] if n<len(source["receipts"]) else None
        val=source["values"][n] if n<len(source["values"]) else None
        result.append({"n":n+1,"date":d["original"] if d else "FALTA", "receipt":r or "FALTA",
                       "value":val["value"] if val else "FALTA", "origin":val.get("origin", "") if val else "",
                       "date_problem": d and d["status"] != "VALIDA"})
    return result

def analyze(path):
    wb=load_workbook(path, data_only=False, read_only=True)
    incidents=[]; stats={"sheets":0,"rows":0,"compatible_sheets":[]}
    for ws in wb.worksheets:
        active=None; context="TABLA PRINCIPAL"; has_historical_words=False; sheet_used=False
        # enumerate is deliberate: xlsx rows keep their Excel index including blanks.
        for rn, row in enumerate(ws.iter_rows(values_only=True), 1):
            values=list(row)
            found={}
            for col, value in enumerate(values, 1):
                kind=header_kind(value)
                if kind: found.setdefault(kind, []).append((col, cell_text(value)))
            if len(found) >= 2 and ("receipts" in found or "fid_receipts" in found):
                active={k:v for k,v in found.items()}; sheet_used=True
                continue
            if not active: continue
            # El texto del encabezado CESIONES/TRASLADOS no convierte toda la hoja
            # en cesión: solo los marcadores hallados en un registro/histórico.
            detected_context=context_marker(values, "TABLA PRINCIPAL")
            row_has_marker = detected_context != "TABLA PRINCIPAL"
            # Un marcador dentro de una fila pertenece a esa fila. Los marcadores
            # aislados siguen actuando como contexto para una tabla posterior.
            row_context = detected_context if row_has_marker else context
            # A new narrative line does not carry payments. Process only rows with relevant cells.
            def get_one(k):
                return cell_text(values[active[k][0][0]-1]) if k in active and active[k][0][0] <= len(values) else ""
            c_dates_raw=split_dates(get_one("dates")); c_receipts=split_receipts(get_one("receipts")); c_values=split_values(get_one("received"))
            f_dates_raw=[x for x in split_dates(get_one("dates")) if x.upper().replace(" ","").endswith("F")]
            f_receipts=split_receipts(get_one("fid_receipts")); f_values=[]
            for col, header in active.get("fid_value", []):
                raw=cell_text(values[col-1]) if col <= len(values) else ""
                for val in split_values(raw): f_values.append({"value":val,"origin":header,"period":fiduciary_period(header),"column":col})
            # Non-F dates belong to constructora. A date column may contain both sources.
            c_dates_raw=[x for x in c_dates_raw if not x.upper().replace(" ","").endswith("F")]
            if not any([c_dates_raw,c_receipts,c_values,f_dates_raw,f_receipts,f_values]):
                if row_has_marker:
                    context=row_context; has_historical_words=True
                continue
            if row_has_marker:
                has_historical_words=True
            else:
                context="TABLA PRINCIPAL"
            stats["rows"]+=1
            evidence={"constructora":{"dates":[asdict(parse_date(x)) for x in c_dates_raw],"receipts":c_receipts,"values":[{"value":x} for x in c_values]},
                      "fiduciaria":{"dates":[asdict(parse_date(x)) for x in f_dates_raw],"receipts":f_receipts,"values":f_values}}
            incidents.extend(gf_blocking(ws.title, rn, row_context, active, values, evidence))
        if sheet_used:
            stats["sheets"]+=1; stats["compatible_sheets"].append(ws.title)
        elif has_historical_words:
            incidents.append(make_incident("STRUCTURE",ws.title,0,"CESIÓN/TRASLADO","Se detectó un marcador histórico pero no encabezados suficientes para analizarlo.",["ENCABEZADOS REQUERIDOS: RECIBOS/FIDUBOGOTA y FECHA"],{}))
    # Nivel 2: una incidencia por fila y tipo. Las evidencias individuales se
    # conservan para no perder trazabilidad ni el detalle.
    grouped = {}
    for inc in incidents:
        key = (inc["kind"], inc["sheet"], inc["row"], inc["context"])
        bucket = grouped.setdefault(key, {"base": inc, "items": [], "seen": set()})
        signature = (tuple(inc["info"]), inc["description"])
        if signature not in bucket["seen"]:
            bucket["seen"].add(signature)
            bucket["items"].append(inc)
    unique=[]
    for bucket in grouped.values():
        inc = bucket["base"].copy()
        items = bucket["items"]
        inc["evidence_items"] = items
        count = len(items)
        inc["info"] = [f"EVIDENCIAS BLOCKING EN ESTA FILA: {count}"]
        if count > 1:
            inc["description"] = f"{inc['description']} Se agruparon {count} evidencias blocking de Gestión Fiduciaria."
        inc["id"] = len(unique)
        unique.append(inc)
    wb.close()
    return {
        "incidents": unique,
        "stats": stats,
        "groups": dict(Counter(x["kind"] for x in unique)),
        "evidence_counts": dict(Counter({kind: sum(len(x["evidence_items"]) for x in unique if x["kind"] == kind) for kind in {x["kind"] for x in unique}})),
    }
